import openpyxl
import pandas as pd

# Public portfolio version: raw source files are intentionally excluded.
# Replace these placeholders with local anonymized/raw inputs if regenerating.
DK_FILE = "data/raw/market_dk_assets.xlsx"
DE_FILE = "data/raw/market_de_assets.xlsx"
OUTPUT = "data/processed/clean_data.csv"
TASK2_OUTPUT = "data/processed/task2_asset_summary.csv"


def _load(path, sheet):
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = list(wb[sheet].iter_rows(min_row=3, values_only=True))
    wb.close()
    df = pd.DataFrame(rows[1:], columns=rows[0])
    df = df.loc[:, df.columns.notna()]
    return df[df["Production reporting period"].notna()].copy()


def _process_dk(df):
    def certainty(row):
        if pd.notna(pd.to_numeric(row["Final production, total share (MWh)"], errors="coerce")):
            return "Final"
        if pd.notna(pd.to_numeric(row["Preliminary production, total share (MWh)"], errors="coerce")):
            return "Preliminary"
        return "Estimated"

    lost   = pd.to_numeric(df["Lost production due to wind (MWh)"],   errors="coerce")
    gained = pd.to_numeric(df["Gained production due to wind (MWh)"], errors="coerce")
    budget = pd.to_numeric(df["Production in budget 2025, Momentum ownership share (MWh)"], errors="coerce")
    actual = pd.to_numeric(df["Actual production, Momentum ownership share (MWh): Output to use"], errors="coerce")
    avail  = pd.to_numeric(df["Availability (incl. long-term stoppage) %"], errors="coerce")   # unit: 0–1
    normalised = pd.to_numeric(df["Normalised Production (wind index corrected), MWh"], errors="coerce")

    deviation          = actual - budget
    wind_impact        = actual - normalised                            # negative = weak wind impact
    wind_adj_deviation = normalised - budget                            # weather-adjusted residual

    return pd.DataFrame({
        "asset_id":                  df["Asset ID"],
        "market":                    df["Country"],
        "price_zone":                df["Price zone"],
        "technology":                df["Technology"],
        "turbine_model":             df["Turbine model"],
        "commissioning_year":        pd.to_datetime(df["Turbine COD date"], errors="coerce").dt.year.astype("Int64"),
        "ownership_budget":          pd.to_numeric(df["Ownership share (in budget)"], errors="coerce"),
        "ownership_actual":          pd.to_numeric(df["Ownership share (actual)"],    errors="coerce"),
        "capacity_mw":               pd.to_numeric(df["Capacity, total (MW)"],        errors="coerce"),
        "reporting_period":          pd.to_datetime(df["Production reporting period"]),
        "budget_mwh":                budget,
        "actual_mwh":                actual,
        "deviation_mwh":             deviation,
        "deviation_pct":             deviation / budget,                # negative = underperformance
        "production_certainty":      df.apply(certainty, axis=1),       # Final / Preliminary / Estimated
        "availability":              avail,                             # 0–1 (fraction)
        "availability_standardised": avail,                             # DK already 0–1
        "normalised_mwh":            normalised,
        "wind_index":                pd.to_numeric(df["Windindex relative to standard"],    errors="coerce"),
        "avg_wind_ms":               pd.to_numeric(df["Average wind speed (m/s)"],          errors="coerce"),
        "lost_production_wind_mwh":  lost,
        "gained_production_wind_mwh": gained,
        "wind_impact_mwh":           wind_impact,
        "wind_adjusted_deviation_mwh": wind_adj_deviation,
        "stoppage_mwh":              pd.to_numeric(df["Long-term stoppage (MWh)"],           errors="coerce"),
        "loss_unavail_mwh":          pd.to_numeric(df["Loss from unavailability (beyond long-term stoppage), MWh"], errors="coerce"),
        "availability_unit":         "fraction",
    })


def _process_de(df):
    def certainty(row):
        # Realised ≈ Final;  Settled ≈ Preliminary  (verified against Output to use)
        if pd.notna(pd.to_numeric(row["Realised Production, total share (MWh)"], errors="coerce")):
            return "Final"
        if pd.notna(pd.to_numeric(row["Settled production, total share (MWh)"], errors="coerce")):
            return "Preliminary"
        return "Estimated"

    lost   = pd.to_numeric(df["Lost production due to wind (MWh)"],   errors="coerce")
    gained = pd.to_numeric(df["Gained production due to wind (MWh)"], errors="coerce")
    budget = pd.to_numeric(df["Production in budget 2025, MEG ownership share (MWh)"], errors="coerce")
    actual = pd.to_numeric(df["Actual production, MEG ownership share (MWh): Output to use"], errors="coerce")
    avail  = pd.to_numeric(df["Availability (incl. long-term stoppage), index"], errors="coerce")  # unit: 0–100
    normalised = pd.to_numeric(df["Normalised Production (wind index corrected), MWh"], errors="coerce")

    deviation          = actual - budget
    wind_impact        = actual - normalised
    wind_adj_deviation = normalised - budget

    return pd.DataFrame({
        "asset_id":                  df["Asset ID"],
        "market":                    df["Country"],
        "price_zone":                None,
        "technology":                df["Technology"],
        "turbine_model":             df["Turbine model"],
        "commissioning_year":        pd.to_datetime(df["Turbine COD date"], errors="coerce").dt.year.astype("Int64"),
        "ownership_budget":          pd.to_numeric(df["Ownership share (in budget)"], errors="coerce"),
        "ownership_actual":          pd.to_numeric(df["Ownership share (actual)"],    errors="coerce"),
        "capacity_mw":               pd.to_numeric(df["Capacity, total (MW)"],        errors="coerce"),
        "reporting_period":          pd.to_datetime(df["Production reporting period"]),
        "budget_mwh":                budget,
        "actual_mwh":                actual,
        "deviation_mwh":             deviation,
        "deviation_pct":             deviation / budget,
        "production_certainty":      df.apply(certainty, axis=1),
        "availability":              avail,                             # 0–100 (index, raw)
        "availability_standardised": avail / 100,                      # normalised to 0–1 for cross-market comparison
        "normalised_mwh":            normalised,
        "wind_index":                pd.to_numeric(df["Windindex relative to standard"],    errors="coerce"),
        "avg_wind_ms":               pd.to_numeric(df["Average wind speed (m/s)"],          errors="coerce"),
        "lost_production_wind_mwh":  lost,
        "gained_production_wind_mwh": gained,
        "wind_impact_mwh":           wind_impact,
        "wind_adjusted_deviation_mwh": wind_adj_deviation,
        "stoppage_mwh":              pd.to_numeric(df["Long-term stoppage (MWh)"],           errors="coerce"),
        "loss_unavail_mwh":          pd.to_numeric(df["Loss from unavailability (beyond long-term stoppage), MWh"], errors="coerce"),
        "availability_unit":         "index_raw",                      # original unit before standardisation
    })


# ── combine ───────────────────────────────────────────────────────────
dk = _process_dk(_load(DK_FILE, "Denmark"))
de = _process_de(_load(DE_FILE, "Germany"))
df = pd.concat([dk, de], ignore_index=True)

# ownership tier (Task 1 case requirement)
df["ownership_tier"] = df["ownership_actual"].apply(
    lambda x: "Majority (>50%)" if pd.notna(x) and x > 0.5 else "Minority (≤50%)"
)

# exclusion reason — transparent flagging, not silent drop
def _flag(row):
    if row["ownership_actual"] == 0 and pd.notna(row["budget_mwh"]) and row["budget_mwh"] > 0:
        return "ownership_zero_budget_nonzero"   # DE-0017 type: divested but budget still present
    if row["ownership_actual"] == 0:
        return "ownership_zero"
    if pd.isna(row["availability_standardised"]):
        return "availability_missing"
    return "ok"

df["exclusion_reason"] = df.apply(_flag, axis=1)
# operational portfolio = exclude zero-ownership anomalies only
# availability_missing rows still have valid production data — keep them in portfolio view
df["in_operational_portfolio"] = ~df["exclusion_reason"].isin([
    "ownership_zero_budget_nonzero",
    "ownership_zero"
])

df.to_csv(OUTPUT, index=False)


def _weighted_average(values, weights):
    valid = values.notna() & weights.notna() & (weights > 0)
    if not valid.any():
        return pd.NA
    return (values[valid] * weights[valid]).sum() / weights[valid].sum()


def _mode_or_na(series):
    mode = series.dropna().mode()
    return mode.iloc[0] if len(mode) else pd.NA


def _primary_driver(row):
    if row["deviation_mwh"] >= 0:
        return "On / above budget"

    abs_dev = abs(row["deviation_mwh"])
    if abs_dev < 100:
        return "Near budget"

    wind_shortfall = max(-row["wind_impact_mwh"], 0)
    residual_shortfall = max(-row["wind_adjusted_deviation_mwh"], 0)
    wind_share = wind_shortfall / abs_dev
    residual_share = residual_shortfall / abs_dev

    if pd.notna(row["avg_availability"]) and row["avg_availability"] < 0.90 and residual_share >= 0.35:
        return "Availability / residual"
    if wind_share >= 0.70 and residual_share < 0.35:
        return "Wind-driven"
    if residual_share >= 0.60 and wind_share < 0.45:
        return "Residual-driven"
    if wind_share >= 0.25 and residual_share >= 0.25:
        return "Mixed wind and residual"
    return "Review required"


def _build_task2_asset_summary(source):
    op = source[source["in_operational_portfolio"]].copy()

    records = []
    group_cols = ["asset_id", "market", "ownership_tier", "technology", "price_zone"]
    for keys, g in op.groupby(group_cols, dropna=False):
        budget = g["budget_mwh"].sum()
        actual = g["actual_mwh"].sum()
        deviation = g["deviation_mwh"].sum()
        normalised = g["normalised_mwh"].sum()
        wind_impact = g["wind_impact_mwh"].sum()
        wind_adjusted = g["wind_adjusted_deviation_mwh"].sum()
        avg_wind_index = _weighted_average(g["wind_index"], g["budget_mwh"])
        avg_availability = _weighted_average(g["availability_standardised"], g["budget_mwh"])

        records.append({
            "asset_id": keys[0],
            "market": keys[1],
            "ownership_tier": keys[2],
            "technology": keys[3],
            "price_zone": keys[4],
            "turbine_model": g["turbine_model"].iloc[0],
            "commissioning_year": g["commissioning_year"].iloc[0],
            "months_reported": g["reporting_period"].nunique(),
            "capacity_mw": g["capacity_mw"].max(),
            "budget_mwh": budget,
            "actual_mwh": actual,
            "normalised_mwh": normalised,
            "deviation_mwh": deviation,
            "deviation_pct": deviation / budget if budget else pd.NA,
            "wind_impact_mwh": wind_impact,
            "wind_adjusted_deviation_mwh": wind_adjusted,
            "wind_impact_share_of_abs_deviation": abs(wind_impact) / abs(deviation) if deviation else pd.NA,
            "avg_wind_index": avg_wind_index,
            "avg_availability": avg_availability,
            "stoppage_mwh": g["stoppage_mwh"].sum(min_count=1),
            "loss_unavail_mwh": g["loss_unavail_mwh"].sum(min_count=1),
            "availability_missing_months": g["availability_standardised"].isna().sum(),
            "wind_data_missing_months": g["wind_index"].isna().sum(),
            "production_certainty_mode": _mode_or_na(g["production_certainty"]),
        })

    summary = pd.DataFrame(records)
    summary["primary_driver"] = summary.apply(_primary_driver, axis=1)

    underperforming = summary[summary["deviation_mwh"] < 0].copy()
    underperforming["underperformance_rank_mwh"] = (
        underperforming["deviation_mwh"].rank(method="dense", ascending=True).astype(int)
    )
    summary = summary.merge(
        underperforming[["asset_id", "underperformance_rank_mwh"]],
        on="asset_id",
        how="left"
    )

    def priority(row):
        if row["deviation_mwh"] >= 0:
            return "Monitor"
        rank = row["underperformance_rank_mwh"]
        if (pd.notna(rank) and rank <= 10) or row["deviation_pct"] <= -0.30 or row["wind_adjusted_deviation_mwh"] <= -500:
            return "High"
        if (pd.notna(rank) and rank <= 20) or row["deviation_pct"] <= -0.15:
            return "Medium"
        return "Low"

    def data_flag(row):
        flags = []
        if row["availability_missing_months"] > 0:
            flags.append("availability_missing")
        if row["wind_data_missing_months"] > 0:
            flags.append("wind_index_missing")
        if row["market"] == "Germany" and pd.notna(row["loss_unavail_mwh"]) and row["loss_unavail_mwh"] > 10000:
            flags.append("unavailability_loss_not_comparable")
        return "ok" if not flags else "; ".join(flags)

    summary["investigation_priority"] = summary.apply(priority, axis=1)
    summary["data_quality_flag"] = summary.apply(data_flag, axis=1)

    ordered_cols = [
        "asset_id", "market", "ownership_tier", "technology", "price_zone",
        "turbine_model", "commissioning_year",
        "months_reported", "capacity_mw", "budget_mwh", "actual_mwh",
        "normalised_mwh", "deviation_mwh", "deviation_pct",
        "wind_impact_mwh", "wind_adjusted_deviation_mwh",
        "wind_impact_share_of_abs_deviation", "avg_wind_index",
        "avg_availability", "stoppage_mwh", "loss_unavail_mwh",
        "availability_missing_months", "wind_data_missing_months",
        "production_certainty_mode", "primary_driver",
        "underperformance_rank_mwh", "investigation_priority", "data_quality_flag",
    ]
    return summary[ordered_cols].sort_values(
        ["investigation_priority", "deviation_mwh"],
        ascending=[True, True]
    )


task2_summary = _build_task2_asset_summary(df)
task2_summary.to_csv(TASK2_OUTPUT, index=False)

# ── sanity check ──────────────────────────────────────────────────────
print(f"\nRows: {len(df)}  (DK={len(dk)}, DE={len(de)})")
print(f"\nexclusion_reason:\n{df['exclusion_reason'].value_counts()}")
print(f"\nproduction_certainty:\n{df['production_certainty'].value_counts()}")
print(f"\nownership_tier:\n{df['ownership_tier'].value_counts()}")
print(f"\nperiods: {df['reporting_period'].min().date()} → {df['reporting_period'].max().date()}")

op = df[df["in_operational_portfolio"]]
print(f"\nOperational portfolio ({len(op)} rows):")
print(f"  Budget:    {op['budget_mwh'].sum()/1000:.2f}K MWh")
print(f"  Actual:    {op['actual_mwh'].sum()/1000:.2f}K MWh")
print(f"  Deviation: {op['deviation_mwh'].sum()/1000:.2f}K MWh")
print(f"  Dev %:     {op['deviation_mwh'].sum()/op['budget_mwh'].sum()*100:.1f}%")
print(f"\nAll rows (incl. anomalies):")
print(f"  Budget:    {df['budget_mwh'].sum()/1000:.2f}K MWh")
print(f"  Actual:    {df['actual_mwh'].sum()/1000:.2f}K MWh")
print(f"\nTask 2 asset summary:")
print(f"  Assets:    {len(task2_summary)}")
print(f"  Output:    {TASK2_OUTPUT}")
print(f"  High priority assets: {(task2_summary['investigation_priority'] == 'High').sum()}")
